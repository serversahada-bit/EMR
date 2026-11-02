# Urai SATU file Progres (.xlsx) -> baris untuk tabel gabungan per kanal.
#   python urai_progres.py "<file.xlsx>" ["<judul>"]  ->  JSON di stdout
#
# Dipakai importer EMR: link Google Sheet diunduh sebagai .xlsx, lalu file
# ini memanggil pengurai yang SAMA dengan yang mengisi database selama ini
# (extract_progres/google/snack/mp, build_all_cos, build_all_crm). Tidak ada
# logika urai baru di sini -- hanya memilih tab & memetakan ke kolom tabel.
#
# Produk, bulan, tahun diambil dari JUDUL ("Progres Oktober Gamamilk 2026"),
# karena label tanggal di dalam sheet sering basi. Baris yang tanggalnya
# bukan bulan itu sudah dibuang oleh tiap pengurai.
#
# Yang TIDAK dilakukan di sini (dikerjakan importer di EMR, yang punya akses
# database): deteksi sheet beku terhadap bulan sebelumnya, pengisian id_adv,
# dan penulisan ke database.
import contextlib
import datetime
import io
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.argv = sys.argv[:1] + ["2026"] + sys.argv[1:]   # modul build_* membaca argv[1]
ARGS = sys.argv[2:]

import build_all_cos as BC          # noqa: E402
import build_all_crm as BR          # noqa: E402
import extract_google as EG         # noqa: E402
import extract_mp as EM             # noqa: E402
import extract_progres as EP        # noqa: E402
import extract_snack as ES          # noqa: E402

BLN_NO = {v: k for k, v in EP.BULAN.items()}
PENDEK = (("premium", "Premium"), ("plus", "Plus"), ("sereal", "Sereal"),
          ("phenobody", "Phenobody"), ("g_nait", "G'Nait"), ("g'nait", "G'Nait"),
          ("gnait", "G'Nait"), ("etacefit", "Etacefit"))
LENGKAP = {"Gamamilk": "Gamamilk", "Premium": "Gamamilk Premium", "Plus": "Gamamilk Plus",
           "Sereal": "Gamamilk Sereal", "Phenobody": "Phenobody", "G'Nait": "G'Nait",
           "Etacefit": "Etacefit"}
SLUG = {"Gamamilk": "utama", "Premium": "premium", "Plus": "plus", "Sereal": "sereal",
        "Phenobody": "phenobody", "G'Nait": "gnait", "Etacefit": "etacefit"}
# nama sumber Meta lama (dipakai EMR sebagai id sumber) -- lihat meta-queries.ts
META_LAMA = {("Gamamilk", 2026): "data_spend_meta_harian", ("Gamamilk", 2025): "data_spend_meta_2025_harian",
             ("Gamamilk", 2024): "data_spend_meta_2024_harian", ("Premium", 2026): "data_spend_meta_premium_harian",
             ("Premium", 2025): "data_spend_meta_premium_2025_harian", ("Plus", 2026): "data_spend_meta_plus_harian",
             ("Sereal", 2026): "data_spend_meta_sereal_harian", ("G'Nait", None): "data_spend_meta_gnait_harian",
             ("Phenobody", None): "data_spend_meta_phenobody_harian", ("Etacefit", None): "data_spend_meta_etacefit_harian"}
ALIAS_PIC = {"MUAMAR": "AMAR", "GUSTI": "GUSTY", "IMAM GM": "IMAM"}


def produk_dari(judul):
    low = judul.lower()
    return next((v for k, v in PENDEK if k in low), "Gamamilk")


def baca_judul(judul):
    bulan = next((b for b in BLN_NO if re.search(r"\b%s\b" % b, judul, re.I)), None)
    th = re.search(r"20\d\d", judul)
    return bulan, (int(th.group(0)) if th else None), produk_dari(judul)


def sumber_tabel(kanal, prod, tahun):
    if kanal == "meta":
        return META_LAMA.get((prod, tahun)) or META_LAMA.get((prod, None)) or \
            "data_spend_meta_%s_%d_harian" % (SLUG[prod], tahun)
    return "data_spend_%s_%s_%d_harian" % (kanal, SLUG[prod], tahun)


def js(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v != v:
        return None
    return v


def bersih(d):
    return {k: js(v) for k, v in d.items() if not k.startswith("_")}


def adalah_google(ws, title):
    return bool(re.search("google", title, re.I)) or any(re.search("google", x, re.I) for x in EG.banner(ws))


def urai(path, judul):
    bulan, tahun, prod = baca_judul(judul)
    if not bulan or not tahun:
        raise SystemExit(json.dumps({"galat": "Judul '%s' tidak memuat bulan dan tahun "
                                     "(contoh: 'Progres Oktober Gamamilk 2026')." % judul}))
    bno, full, berkas = BLN_NO[bulan], LENGKAP[prod], judul + ".xlsx"
    umum = {"produk": full, "tahun": tahun, "bulan": bulan, "file_sumber": berkas}
    out = {k: [] for k in ("meta", "google", "snack", "mp", "cos_harian", "cos_bulanan",
                           "crm_blast", "crm_ro", "crm_cs")}
    dipakai, dilewati = {}, {}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    judul_tab = list(wb.sheetnames)

    # ---------------- Meta / Google / Snack: sheet ADV ----------------
    meta_per_sheet = {}
    for t in judul_tab:
        u = t.strip().upper()
        if not (u.startswith("ADV ") or u.startswith("COPY OF ADV ")) or "TIKTOK" in u:
            continue
        ws = wb[t]
        try:
            if "SNACK" in u:
                if u == "ADV TESTING SNACK":
                    dilewati[t] = "layout Snack lain (tanpa kolom Spend Iklan)"
                    continue
                recs, _, _ = ES.extract_ws(ws, tahun, bno)
                pic = re.sub(r"\(SNACK\)|SNACK", "", u.replace("COPY OF ", "").replace("ADV", "", 1)).strip()
                kanal = "snack"
            elif adalah_google(ws, t):
                recs, _ = EG.extract_ws(ws, tahun, bno)
                pic = re.sub(r"^(COPY OF\s+)?ADV\s+(GOOGLE\s+)?", "", u).strip()
                kanal = "google"
            else:
                recs, _, _ = EP.extract_ws(ws, tahun, bno)
                pic = re.sub(r"^(COPY OF\s+)?ADV\s+", "", u).replace("META ", "").strip()
                kanal = "meta"
        except Exception as e:                                   # sheet rusak/aneh
            dilewati[t] = "gagal diurai: %s" % str(e)[:80]
            continue
        if not recs:
            dilewati[t] = "tidak ada tanggal %s %d" % (bulan, tahun)
            continue
        pic = ALIAS_PIC.get(pic, pic)
        baris = [dict(bersih(r), **umum, pic=pic, sheet_sumber=t,
                      sumber_tabel=sumber_tabel(kanal, prod, tahun)) for r in recs]
        meta_per_sheet[t] = (kanal, baris)

    # "Copy of ADV X" hanya dipakai bila "ADV X" tidak punya data bulan ini
    for t, (kanal, baris) in meta_per_sheet.items():
        if t.upper().startswith("COPY OF "):
            asli = t[len("Copy of "):]
            if asli in meta_per_sheet:
                dilewati[t] = "salinan dari '%s' yang sudah dipakai" % asli
                continue
        out[kanal].extend(baris)
        dipakai[t] = (kanal, len(baris))

    # ---------------- Marketplace ----------------
    resmi = [x.upper() for x in EM.sheet_resmi(tahun, bno)]
    for t in judul_tab:
        if t.strip().upper() not in resmi:
            continue
        blok, _ = EM.blok_ws(wb[t], tahun, bno, bulan_basi=True)
        kanal = "TIKTOK" if "TIKTOK" in t.upper() else "SHOPEE/LAZADA"
        n = 0
        for bl in blok:
            for r in bl["baris"]:
                rec = {f: js(r.get(f)) for f in EM.FIELDS}
                rec.update(umum, kanal=kanal, agregat=1 if bl["agregat"] else 0, layout=bl["layout"],
                           blok=(bl["label"] or "")[:80], sheet_sumber=t, tanggal=js(r["tanggal"]),
                           label_basi=bl.get("label_basi"),
                           sumber_tabel=sumber_tabel("mp", prod, tahun))
                out["mp"].append(rec)
                n += 1
        if n:
            dipakai[t] = ("mp", n)
        else:
            dilewati[t] = "tidak ada tanggal %s %d" % (bulan, tahun)
    wb.close()

    # ---------------- COS (keluarga Gamamilk) ----------------
    if "gamamilk" in judul.lower():
        semua, ringkas, log = [], [], []
        with contextlib.redirect_stdout(io.StringIO()):
            BC.olah_file(path, berkas, bulan, prod, tahun, {}, semua, ringkas, log)
        for _, bn, bl, jenis, r, _ in semua:
            sp, om = r.get("spend_total"), r.get("omset_total")
            rec = {f: js(r.get(f)) for f in BC.FIELDS}
            rec.update(umum, bulan_no=bn, jenis=jenis, tanggal=js(r["tanggal"]), hari_ke=r["hari_ke"],
                       label_tanggal=js(r["label"]), cos_hitung=(sp / om) if sp and om else None,
                       keterangan=(r.get("keterangan") or None) and str(r["keterangan"])[:255],
                       sumber_tabel="data_cos_%d_harian" % tahun)
            out["cos_harian"].append(rec)
        for x in ringkas:
            (_, bl, bn, hari, sp, om, cos, kor, t_sp, t_om, t_cos, s_sp, s_om, rumus, basi, salin) = x
            out["cos_bulanan"].append(dict(umum, bulan_no=bn, hari=hari, spend_harian=sp, omset_harian=om,
                                           cos_harian=cos, koreksi_total=kor, spend_sheet=t_sp, omset_sheet=t_om,
                                           cos_sheet=t_cos, sisa_selisih_spend=s_sp, sisa_selisih_omset=s_om,
                                           rumus_koreksi=(rumus or None) and rumus[:500], label_basi=1 if basi else 0,
                                           salinan=1 if salin else 0, sumber_tabel="data_cos_%d_bulanan" % tahun))
        if out["cos_harian"]:
            dipakai["COST OF SALES"] = ("cos", len(out["cos_harian"]))
        for l in log:
            if l[5] != "OK" and "COST OF SALES" not in dipakai:
                dilewati["COST OF SALES"] = l[5]

    # ---------------- CRM & blast ----------------
    hasil, log = {"BLAST": [], "CRM": [], "CS": []}, []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    with contextlib.redirect_stdout(io.StringIO()):
        import collections
        BR.olah_wb(wb, berkas, bulan, prod, tahun, {}, hasil, log, collections.Counter())
    wb.close()
    for r in hasil["BLAST"]:
        out["crm_blast"].append(dict({f: js(r.get(f)) for f in BR.F_BLAST}, **umum, tanggal=js(r["tanggal"]),
                                     label_tanggal=js(r["label"]), sheet_sumber=r["sheet"]))
    for r in hasil["CRM"]:
        out["crm_ro"].append(dict({f: js(r.get(f)) for f in BR.F_CRM}, **umum, tanggal=js(r["tanggal"]),
                                  cs=r["cs"], segmen=r["segmen"], label_tanggal=js(r["label"]), sheet_sumber=r["sheet"]))
    for r in hasil["CS"]:
        if not any(r.get(f) for f in BR.F_CS):
            continue
        out["crm_cs"].append(dict({f: js(r.get(f)) for f in BR.F_CS}, **umum, tanggal=js(r["tanggal"]),
                                  jenis_cs=r["jenis_cs"], cs=r["cs"], blok=r["blok"], sumber=r["sumber"],
                                  label_tanggal=js(r["label"]), sheet_sumber=r["sheet"]))
    for l in log:
        (_, _, _, t, jenis, status, cat) = l
        if status == "DIAMBIL":
            dipakai[t] = ("crm_" + {"BLAST": "blast", "CRM": "ro", "CS": "cs"}[jenis], int(cat.split(" ")[0]))
        else:
            dilewati[t] = cat.replace("DILEWATI: ", "")

    for t in judul_tab:
        if t not in dipakai and t not in dilewati:
            dilewati[t] = "bukan sheet yang diimpor"
    return {
        "judul": judul, "produk": full, "produk_pendek": prod, "tahun": tahun, "bulan": bulan, "bulan_no": bno,
        "tabel": out,
        "tab_dipakai": [{"tab": t, "kanal": k, "baris": n} for t, (k, n) in dipakai.items()],
        "tab_dilewati": [{"tab": t, "alasan": a} for t, a in dilewati.items()],
    }


if __name__ == "__main__":
    if not ARGS:
        raise SystemExit("pakai: python urai_progres.py <file.xlsx> [judul]")
    path = ARGS[0]
    judul = ARGS[1] if len(ARGS) > 1 else os.path.splitext(os.path.basename(path))[0]
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(urai(path, judul), ensure_ascii=False, default=str))
