# Extractor sheet iklan SNACK VIDEO dari file Progres.
# Header acuan: "Progres Oktober Gamamilk 2026.xlsx" sheet "ADV PUTRA (SNACK)",
# baris 3 (grup) + baris 4 (sub), kolom c1..c48.
#
# Label kolom yang sama ditulis beda-beda antar sheet/bulan -- mis.
# "FORM IKLAN" vs "Form Dashboard", "FORM OO" vs "Form Order Online",
# "LR FORM LOOPS" vs "LR Form Online" -- dan posisinya bergeser. Jadi
# pemetaan berbasis label + alias, bukan posisi, dan fungsi
# label_tak_termap() dipakai untuk memeriksa apakah ada label baru yang
# belum dikenali (jangan percaya "jumlah kolom termap" sebagai tanda sehat).
import re
import datetime
import openpyxl

from extract_progres import cell, BULAN

BATAS_KOLOM = 52

# (section, pola label) -> nama kanonik. Pola dicocokkan dgn label
# ternormalisasi; lihat pick() untuk urutan exact -> startswith -> substring.
RULES = {
    "BASE": [("date", "tanggal"),
             ("perencanaan target", "target_spend_rencana"),
             ("spend iklan + ppn", "spend_iklan_ppn"),
             ("spend iklan", "spend_iklan")],
    "AKTUAL": [("spanding aktual", "spend_aktual"),
               ("spending aktual", "spend_aktual"),
               ("spend aktual", "spend_aktual")],
    "IKLAN": [("impresi", "impresi"), ("cpm", "cpm"), ("ctr", "ctr"),
              ("klik tautan", "klik_tautan"), ("tayang landas", "tayang_landas"),
              ("page view", "page_view"), ("pageview", "page_view"),
              ("rasio konversi", "rasio_konversi"), ("cpc", "cpc_total")],
    "DASBOARD": [("form iklan", "lead_form_iklan"),
                 ("form dashboard", "lead_form_iklan"),
                 ("form dasboard", "lead_form_iklan"),
                 ("form oo", "lead_form_oo"),
                 ("form order online", "lead_form_oo"),
                 ("lr form", "lr_form"),
                 ("wa iklan", "lead_wa_iklan"),
                 ("total lead dasboard", "total_lead_dashboard"),
                 ("total lead dashboard", "total_lead_dashboard"),
                 ("cost per result", "cpr"),
                 ("cpa dashboard", "cpr")],
    "REAL": [("form", "lead_real_form"),
             ("wa", "lead_real_wa"),
             ("lead wa abandoned", "lead_wa_abandoned"),
             ("wa abandoned", "lead_wa_abandoned"),
             ("lead hilang", "lead_hilang"),
             ("total lead real", "total_lead_real"),
             ("cpl", "cpl"),
             ("target lead", "target_lead")],
    "RATE": [("form", "lead_rate_form"), ("wa", "lead_rate_wa")],
    "NEW": [("closing customer", "closing_customer_baru"), ("box", "box_baru"),
            ("closing rate", "closing_rate_baru"), ("up selling", "up_selling_baru")],
    "FU": [("closing customer", "closing_customer_fu"), ("box", "box_fu")],
    "TOTAL": [("closing total", "closing_total"), ("box total", "box_total"),
              ("target box", "target_box"), ("closing rate", "closing_rate_total"),
              ("up selling", "up_selling_total"),
              ("biaya akuisi lead", "biaya_akuisisi_lead"),
              ("biaya akuisisi lead", "biaya_akuisisi_lead"),
              ("biaya akuisisi customer", "biaya_akuisisi_customer"),
              ("akuisisi box aktual", "akuisisi_box_aktual"),
              ("biaya akuisisi box", "biaya_akuisisi_box"),
              # khas sheet GUSTI -- dia memakai ini sbg ganti cpl_aktual
              ("cpa dashboard", "cpa_dashboard_non_tax"),
              ("cpa after compensation", "cpa_after_compensation")],
    "TAIL": [("batas akuisisi box", "batas_akuisisi_box_nilai"),
             ("batas", "batas_akuisisi_box_nilai"),
             ("% akusisi box (after compensation)", "persen_akuisisi_box_after_comp"),
             ("% akusisi box", "persen_akuisisi_box"),
             ("% akuisisi box", "persen_akuisisi_box"),
             ("akusisi box", "persen_akuisisi_box"),
             ("akuisisi box", "persen_akuisisi_box"),
             ("grade", "grade"), ("arus", "arus"),
             ("cpl aktual", "cpl_aktual"),
             ("cpa after compensation", "cpa_after_compensation"),
             ("% akusisi box (after", "persen_akuisisi_box_after_comp"),
             ("% akuisisi box (after", "persen_akuisisi_box_after_comp")],
}

GRUP2SEC = {
    "data iklan snack": "IKLAN",
    "lead dasboard": "DASBOARD", "lead dashboard": "DASBOARD",
    "lead real cs": "REAL", "% lead rate cs": "RATE",
    "new customer real hari ini": "NEW", "follow up": "FU",
    "total performa cs": "TOTAL",
}

FIELDS = ["target_spend_rencana", "spend_iklan", "spend_iklan_ppn", "spend_aktual",
          "impresi", "cpm", "ctr", "klik_tautan", "tayang_landas", "page_view",
          "rasio_konversi", "cpc_total",
          "lead_form_iklan", "lead_form_oo", "lr_form", "lead_wa_iklan",
          "total_lead_dashboard", "cpr",
          "lead_real_form", "lead_real_wa", "lead_wa_abandoned",
          "lead_hilang", "total_lead_real",
          "cpl", "target_lead", "lead_rate_form", "lead_rate_wa",
          "closing_customer_baru", "box_baru", "closing_rate_baru", "up_selling_baru",
          "closing_customer_fu", "box_fu",
          "closing_total", "box_total", "target_box", "closing_rate_total",
          "up_selling_total", "biaya_akuisisi_lead", "biaya_akuisisi_customer",
          "akuisisi_box_aktual", "biaya_akuisisi_box", "persen_akuisisi_box",
          "batas_akuisisi_box_nilai", "grade", "arus", "cpl_aktual",
          "cpa_dashboard_non_tax", "cpa_after_compensation",
          "persen_akuisisi_box_after_comp"]

TEKS = {"grade"}


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).replace("\n", " ")).strip().lower()


def pick(section, label):
    """exact -> startswith -> substring; yang terpanjang menang."""
    for tier in (0, 1, 2):
        best = None
        for pat, canon in RULES.get(section, []):
            ok = (label == pat) if tier == 0 else (
                label.startswith(pat) if tier == 1 else (pat in label))
            if ok and (best is None or len(pat) > best[0]):
                best = (len(pat), canon)
        if best:
            return best[1]
    return None


def grp_section(g):
    g = g.rstrip(" %").strip()
    for k, sec in GRUP2SEC.items():
        if g.startswith(k.rstrip(" %")):
            return sec
    if "biaya" not in g and ("akusisi box" in g or "akuisisi box" in g):
        return "TAIL"
    if g.startswith("batas") or g in ("grade", "arus"):
        return "TAIL"
    return None


def _peta(r3, r4):
    """-> (map, daftar label yg tidak termap)"""
    out, section, seen, lepas = {}, "BASE", set(), []
    for j in range(min(BATAS_KOLOM, max(len(r3), len(r4)))):
        g = norm(r3[j]) if j < len(r3) else ""
        s = norm(r4[j]) if j < len(r4) else ""
        if g:
            sec = grp_section(g)
            if sec:
                section = sec
        label = s or g
        if not label:
            continue
        # "Spanding/Spend Aktual" berdiri di antara BASE dan blok iklan
        if "aktual" in label and "akuisisi" not in label and "cpl" not in label:
            canon = pick("AKTUAL", label)
        else:
            canon = pick(section, label)
            if canon is None and section == "BASE":
                canon = pick("IKLAN", label)
            if canon is None and section in ("NEW", "FU"):
                canon = pick("TOTAL", label)
            if canon is None:
                canon = pick("TAIL", label)
        if canon and canon not in seen:
            out[canon] = j
            seen.add(canon)
        elif canon is None:
            lepas.append((j + 1, label))
    return out, lepas


def label_tak_termap(ws):
    """Label header yang belum dikenali -- untuk memeriksa sheet baru."""
    ws.reset_dimensions()
    raw = list(ws.iter_rows(min_row=1, max_row=5, max_col=BATAS_KOLOM, values_only=True))
    if len(raw) < 4:
        return []
    return _peta(list(raw[2]), list(raw[3]))[1]


def banner(ws):
    ws.reset_dimensions()
    raw = list(ws.iter_rows(min_row=1, max_row=2, max_col=BATAS_KOLOM, values_only=True))
    return [str(v).strip() for r in raw for v in r
            if isinstance(v, str) and len(str(v).strip()) > 3]


def extract_ws(ws, tahun, bulan_no, max_scan=400):
    ws.reset_dimensions()
    raw = list(ws.iter_rows(min_row=1, max_row=max_scan, max_col=60, values_only=True))
    if len(raw) < 5:
        return [], {}, []
    cmap, lepas = _peta(list(raw[2]), list(raw[3]))
    if "tanggal" not in cmap or "spend_iklan" not in cmap:
        return [], cmap, lepas
    i_date = cmap["tanggal"]
    rows = []
    last_i = last_d = None
    for off, r in enumerate(raw[4:]):
        i = off + 5
        d = r[i_date] if i_date < len(r) else None
        if not isinstance(d, datetime.datetime):
            continue
        if last_d is not None and (d.date() <= last_d or i - last_i > 3):
            break
        if d.year != tahun or d.month != bulan_no:
            if last_d is None:
                continue
            break
        last_i, last_d = i, d.date()
        rec = {"tanggal": d.date()}
        for f in FIELDS:
            j = cmap.get(f)
            rec[f] = cell(r[j]) if (j is not None and j < len(r)) else None
        sp = rec["spend_iklan"]
        if not (isinstance(sp, (int, float)) and not isinstance(sp, bool) and sp != 0):
            continue
        rows.append(rec)
    return rows, cmap, lepas


def extract(path, sheet, tahun, bulan_no, max_scan=400):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return extract_ws(wb[sheet], tahun, bulan_no, max_scan)
    finally:
        wb.close()
