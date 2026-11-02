# Extractor sheet iklan GOOGLE dari file Progres.
# Header acuan: "Progres Oktober Gamamilk 2026.xlsx" sheet "ADV Ryang",
# baris 3 (grup) + baris 4 (sub), kolom c1..c39.
#
# PENTING: kolom c40 ke kanan adalah TABEL LAIN (banner "PRODUK BARU" di
# baris 1), jadi pemetaan dipotong di c39. Logika blok tanggal & aturan
# baris terisi sama dengan extract_progres.py -- lihat catatan di sana.
import re
import datetime
import openpyxl

from extract_progres import cell, BULAN

BATAS_KOLOM = 39  # c40+ milik tabel lain

# label kanonik -> nama snake. Kuncinya label ternormalisasi (lowercase,
# spasi dirapatkan); sebagian sheet menulisnya sedikit berbeda, mis.
# "Impresi" vs "Tayangan" dan "Page View" vs "Pageview".
LABEL = {
    "date": "tanggal",
    "target spend": "target_spend",
    "spend iklan": "spend_iklan",
    "spend iklan plus pajak": "spend_iklan_ppn",
    "jumlah lead dasboard": "total_lead_dashboard",
    "jumlah lead dashboard": "total_lead_dashboard",
    "jumlah lead real cs": "total_lead_real",
    "jumlah lead real": "total_lead_real",
    "klik tautan non- gsn": "klik_tautan_non_gsn",
    "klik tautan non-gsn": "klik_tautan_non_gsn",
    # Sep/Okt 2026 Plus menulis kolom non-GSN cuma "Klik Tautan",
    # posisinya tetap bersebelahan dengan "Klik Tautan GSN".
    "klik tautan": "klik_tautan_non_gsn",
    "klik tautan gsn": "klik_tautan_gsn",
    "impresi": "impresi",
    "tayangan": "impresi",
    "page view": "page_view",
    "pageview": "page_view",
    "rasio pageview": "rasio_pageview",
    "rasio convert": "rasio_convert",
    "cpc non-gsn": "cpc_non_gsn",
    "cpc non- gsn": "cpc_non_gsn",
    "cpc": "cpc_non_gsn",  # idem: bersebelahan dengan "CPC GSN"
    "cpc gsn": "cpc_gsn",
    "cpc total": "cpc_total",
    "ctr": "ctr",
    "closing new customer": "closing_customer_baru",
    "box new customer": "box_baru",
    "box/pcs new customer": "box_baru",
    "closing follow up": "closing_customer_fu",
    "box/pcs follow up": "box_fu",
    "box follow up": "box_fu",
    "total closing": "closing_total",
    "total box": "box_total",
    "target box": "target_box",
    "total upselling": "up_selling_total",
    "biaya akuisisi lead dashboard": "biaya_akuisisi_lead_dashboard",
    "biaya akuisisi lead dashboard (cpr)": "biaya_akuisisi_lead_dashboard",
    "cpr": "biaya_akuisisi_lead_dashboard",
    "total closing rate": "closing_rate_total",
    "biaya akuisisi lead (cpl)": "cpl",
    "biaya akuisisi customer": "biaya_akuisisi_customer",
    "biaya akuisisi box/pcs": "biaya_akuisisi_box",
    "grade": "grade",
    "profit": "profit",
}

# sublabel FORM/WA dibedakan lewat grupnya
GRUP_FORM_WA = {
    "lead dashboard": ("lead_dashboard_form", "lead_dashboard_wa"),
    "lead real cs": ("lead_real_form", "lead_real_wa"),
    "lead real": ("lead_real_form", "lead_real_wa"),
    "lead rate cs": ("lead_rate_form", "lead_rate_wa"),
    "lead rate": ("lead_rate_form", "lead_rate_wa"),
}

FIELDS = [
    "target_spend", "spend_iklan", "spend_iklan_ppn",
    "lead_dashboard_form", "lead_dashboard_wa", "total_lead_dashboard",
    "lead_real_form", "lead_real_wa", "total_lead_real",
    "lead_rate_form", "lead_rate_wa",
    "klik_tautan_non_gsn", "klik_tautan_gsn", "impresi", "page_view",
    "rasio_pageview", "rasio_convert", "cpc_non_gsn", "cpc_gsn", "cpc_total", "ctr",
    "closing_customer_baru", "box_baru", "closing_customer_fu", "box_fu",
    "closing_total", "box_total", "target_box", "up_selling_total",
    "biaya_akuisisi_lead_dashboard", "closing_rate_total", "cpl",
    "biaya_akuisisi_customer", "biaya_akuisisi_box",
    "batas_akuisisi_box_persen", "batas_akuisisi_box_nilai", "grade", "profit",
]


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).replace("\n", " ")).strip().lower()


def build_map(r3, r4):
    """-> {nama_snake: index kolom 0-based}, hanya c1..c39."""
    out, grup = {}, ""
    for j in range(min(BATAS_KOLOM, max(len(r3), len(r4)))):
        g = norm(r3[j]) if j < len(r3) else ""
        s = norm(r4[j]) if j < len(r4) else ""
        if g:
            grup = g
        canon = None
        if s in ("form", "wa", "form (%)", "wa (%)"):
            pasangan = GRUP_FORM_WA.get(grup)
            if pasangan:
                canon = pasangan[0] if s.startswith("form") else pasangan[1]
        if canon is None:
            label = s or g
            canon = LABEL.get(label)
            if canon is None and label.startswith("batas akuisisi box"):
                # dua kolom berurutan: "(Maksimal %)" lalu "(Dalam Nilai)"
                canon = ("batas_akuisisi_box_nilai"
                         if "nilai" in label
                         else "batas_akuisisi_box_persen")
        if canon and canon not in out:
            out[canon] = j
    return out


def extract_ws(ws, tahun, bulan_no, max_scan=400):
    ws.reset_dimensions()
    raw = list(ws.iter_rows(min_row=1, max_row=max_scan, max_col=60, values_only=True))
    if len(raw) < 5:
        return [], {}
    cmap = build_map(list(raw[2]), list(raw[3]))
    if "tanggal" not in cmap or "spend_iklan" not in cmap:
        return [], cmap
    i_date = cmap["tanggal"]
    rows = []
    last_i = last_d = None
    for off, r in enumerate(raw[4:]):
        i = off + 5
        d = r[i_date] if i_date < len(r) else None
        if not isinstance(d, datetime.datetime):
            continue
        if last_d is not None and (d.date() <= last_d or i - last_i > 3):
            break
        if d.year != tahun or d.month != bulan_no:
            if last_d is None:
                continue
            break
        last_i, last_d = i, d.date()
        rec = {"tanggal": d.date()}
        for f in FIELDS:
            j = cmap.get(f)
            rec[f] = cell(r[j]) if (j is not None and j < len(r)) else None
        sp = rec["spend_iklan"]
        if not (isinstance(sp, (int, float)) and not isinstance(sp, bool) and sp != 0):
            continue
        rows.append(rec)
    return rows, cmap


def extract(path, sheet, tahun, bulan_no, max_scan=400):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return extract_ws(wb[sheet], tahun, bulan_no, max_scan)
    finally:
        wb.close()


def banner(ws):
    """Teks di baris 1-2; dipakai mengenali sheet mana yang Google Ads."""
    ws.reset_dimensions()
    raw = list(ws.iter_rows(min_row=1, max_row=2, max_col=BATAS_KOLOM, values_only=True))
    teks = [str(v).strip() for r in raw for v in r
            if isinstance(v, str) and len(str(v).strip()) > 3]
    return teks
